"""Builds the complete dataset used by this project:

  data/purchase_orders.xlsx
  data/vendors.xlsx
  invoices/*.pdf              (real provided dataset + supplementary synthetic PDFs)
  invoices/ground_truth.json  (expected decision for the integration test)

This repo ships with a REAL provided invoice/PO dataset at
`invoice_po_matching_dataset/` (12 invoice PDFs + a workbook of purchase
orders / vendor master / ground truth -- see that folder's own README).
This script merges that real dataset's PO and vendor facts into
data/purchase_orders.xlsx and data/vendors.xlsx (in this project's own
column schema -- see README section "PO spreadsheet"), copies its invoice
PDFs into invoices/ UNCHANGED, and additionally generates a small set of
supplementary synthetic invoices/POs for edge cases the real dataset does
not cover (split/partial invoicing across multiple invoices, PO-number
formatting variants, a malformed PO reference, bundled line items). The
supplementary POs are numbered PO2xxx to avoid colliding with the real
dataset's PO1xxx / PO9998 numbers.

Run once: `python scripts/generate_dataset.py`. Safe to re-run (overwrites).
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import fitz  # PyMuPDF, used only to compose the scanned/image-only PDF
import openpyxl
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

REPO_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = REPO_ROOT / "data"
INVOICES_DIR = REPO_ROOT / "invoices"
REAL_DATASET_DIR = REPO_ROOT / "invoice_po_matching_dataset"
REAL_WORKBOOK = REAL_DATASET_DIR / "PO_match_dataset.xlsx"


# --------------------------------------------------------------------------
# PO + vendor master data
# --------------------------------------------------------------------------

VENDORS = [
    dict(vendor_id="V001", vendor_name="Apex Office Supplies Pvt Ltd", status="APPROVED", tolerance_type="percentage", tolerance_value=1),
    dict(vendor_id="V002", vendor_name="BrightTech Solutions Pvt Ltd", status="APPROVED", tolerance_type="percentage", tolerance_value=2),
    dict(vendor_id="V003", vendor_name="CloudNine Services Pvt Ltd", status="APPROVED", tolerance_type="percentage", tolerance_value=1),
    # Vendor-level tolerance is the DEFAULT used only when a PO doesn't carry its
    # own override. Delta's real-dataset POs (see load_real_po_rows) specify
    # 2%/1% overrides per PO; this absolute-500 value only applies to Delta POs
    # that don't (e.g. this project's own synthetic PO2007/PO2010).
    dict(vendor_id="V004", vendor_name="Delta Facilities Management", status="APPROVED", tolerance_type="absolute", tolerance_value=500),
    dict(vendor_id="V005", vendor_name="Unknown Vendor", status="NOT_APPROVED", tolerance_type="percentage", tolerance_value=0),
    dict(vendor_id="V006", vendor_name="Evergreen Industrial Supplies", status="APPROVED", tolerance_type="percentage", tolerance_value=1),
    dict(vendor_id="V007", vendor_name="NORTHSTAR SUPPLY CO.", status="NOT_APPROVED", tolerance_type="percentage", tolerance_value=0),
]

# Each entry: PO Number, Vendor Name, PO Date, [(Item, Description, Qty, Unit Price)]
PURCHASE_ORDERS = [
    ("PO2001", "Apex Office Supplies Pvt Ltd", "2026-01-01", [("A4 Paper Ream", "A4 Paper Ream 500 sheets", 100, 250.00)]),
    ("PO2002", "BrightTech Solutions Pvt Ltd", "2026-01-02", [("Laptop Docking Station", "USB-C Laptop Docking Station", 20, 4500.00)]),
    ("PO2003", "CloudNine Services Pvt Ltd", "2026-01-03", [("Cloud Storage Subscription", "Cloud Storage Subscription - Monthly", 12, 2000.00)]),
    ("PO2004", "Apex Office Supplies Pvt Ltd", "2026-01-04", [("Whiteboard Markers Box", "Whiteboard Markers Box", 50, 150.00)]),
    ("PO2005", "BrightTech Solutions Pvt Ltd", "2026-01-05", [("Network Switch 24-port", "Managed Network Switch 24-port", 5, 12000.00)]),
    ("PO2006", "CloudNine Services Pvt Ltd", "2026-01-06", [("Onsite Support Hours", "Onsite Technical Support Hours", 40, 1200.00)]),
    ("PO2007", "Delta Facilities Management", "2026-01-07", [("Facility Maintenance Contract", "Facility Maintenance Contract - Quarterly", 1, 65000.00)]),
    ("PO2008", "Apex Office Supplies Pvt Ltd", "2026-01-08", [("Ergonomic Chair", "Ergonomic Office Chair", 10, 8000.00)]),
    ("PO2009", "BrightTech Solutions Pvt Ltd", "2026-01-09", [("Server Rack Unit", "42U Server Rack Unit", 3, 40000.00)]),
    ("PO2010", "Delta Facilities Management", "2026-01-10", [("Housekeeping Service - Man Days", "Housekeeping Service - Man Days", 100, 1000.00)]),
    ("PO2011", "BrightTech Solutions Pvt Ltd", "2026-01-11", [("USB-C Hub", "USB-C Hub 7-in-1", 25, 800.00)]),
    ("PO2012", "Apex Office Supplies Pvt Ltd", "2026-01-12", [("Pens Box", "Ball Pens Box of 50", 200, 20.00), ("Notebooks Pack", "A5 Notebooks Pack of 10", 100, 50.00)]),
]


def load_real_po_rows() -> list[tuple]:
    """Reads the real dataset's `purchase_orders` sheet and returns rows in
    this project's own column order, with a PO-level percentage tolerance
    override (the real dataset defines tolerance per PO, e.g. Delta's
    PO1004 is 2% while its PO1006/PO1013 are 1%).
    """
    wb = openpyxl.load_workbook(REAL_WORKBOOK, data_only=True)
    ws = wb["purchase_orders"]
    rows = list(ws.iter_rows(values_only=True))
    header, data_rows = rows[0], rows[1:]
    idx = {name: i for i, name in enumerate(header)}

    out = []
    for row in data_rows:
        po_number = row[idx["po_number"]]
        vendor = row[idx["vendor"]]
        description = row[idx["description"]]
        qty = row[idx["quantity"]]
        unit_price = row[idx["unit_price"]]
        amount = row[idx["line_amount"]]
        currency = row[idx["currency"]]
        po_status = row[idx["po_status"]]
        tolerance_pct = row[idx["amount_tolerance_pct"]]
        out.append(
            (
                po_number, vendor, None, description, description, qty, unit_price, amount,
                currency, po_status, "percentage", round(tolerance_pct * 100, 4) if tolerance_pct is not None else None,
            )
        )
    return out


def load_real_vendor_rows() -> list[tuple]:
    wb = openpyxl.load_workbook(REAL_WORKBOOK, data_only=True)
    ws = wb["vendor_master"]
    rows = list(ws.iter_rows(values_only=True))
    header, data_rows = rows[0], rows[1:]
    idx = {name: i for i, name in enumerate(header)}
    out = []
    for row in data_rows:
        vendor = row[idx["vendor"]]
        approval_status = row[idx["approval_status"]]
        status = "APPROVED" if str(approval_status).strip().lower() == "approved" else "NOT_APPROVED"
        out.append((vendor, status))
    return out


def build_vendor_xlsx() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vendors"
    ws.append(["vendor_id", "vendor_name", "status", "tolerance_type", "tolerance_value"])
    known_names = set()
    next_id = 1
    for v in VENDORS:
        ws.append([v["vendor_id"], v["vendor_name"], v["status"], v["tolerance_type"], v["tolerance_value"]])
        known_names.add(v["vendor_name"])
        next_id += 1

    # Merge in any real-dataset vendor not already covered above (status is
    # authoritative from the real vendor_master; default 1% tolerance as a
    # fallback for any PO of theirs that doesn't carry its own override).
    for vendor_name, status in load_real_vendor_rows():
        if vendor_name in known_names:
            continue
        ws.append([f"V{100 + next_id}", vendor_name, status, "percentage", 1 if status == "APPROVED" else 0])
        known_names.add(vendor_name)
        next_id += 1

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(DATA_DIR / "vendors.xlsx")


def build_po_xlsx() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PurchaseOrders"
    ws.append(
        ["PO Number", "Vendor Name", "PO Date", "Item", "Description", "Quantity", "Unit Price", "PO Amount",
         "Currency", "Status", "Tolerance Type", "Tolerance Value"]
    )
    for po_number, vendor_name, po_date, items in PURCHASE_ORDERS:
        for item_name, description, qty, unit_price in items:
            amount = round(qty * unit_price, 2)
            ws.append([po_number, vendor_name, po_date, item_name, description, qty, unit_price, amount, "INR", "OPEN", None, None])

    for row in load_real_po_rows():
        ws.append(list(row))

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(DATA_DIR / "purchase_orders.xlsx")


def copy_real_invoices_and_merge_ground_truth(ground_truth: dict) -> None:
    """Copies the real dataset's invoice PDFs into invoices/ UNCHANGED and
    merges their expected decisions (from invoice_ground_truth) into the
    same ground_truth.json used by the integration test.
    """
    INVOICES_DIR.mkdir(parents=True, exist_ok=True)
    src_dir = REAL_DATASET_DIR / "invoices"
    for pdf_path in src_dir.glob("*.pdf"):
        shutil.copyfile(pdf_path, INVOICES_DIR / pdf_path.name)

    wb = openpyxl.load_workbook(REAL_WORKBOOK, data_only=True)
    ws = wb["invoice_ground_truth"]
    rows = list(ws.iter_rows(values_only=True))
    header, data_rows = rows[0], rows[1:]
    idx = {name: i for i, name in enumerate(header)}

    for row in data_rows:
        filename = row[idx["invoice_file"]]
        expected_decision = row[idx["expected_decision"]]
        edge_case = row[idx["edge_case"]]
        is_scanned = filename.startswith("scanned_")
        entry = {"expected_decision": expected_decision, "notes": f"(real provided dataset) {edge_case}", "source": "invoice_po_matching_dataset"}
        if is_scanned:
            entry["environment_dependent"] = True
            entry["notes"] += " -- scanned/image PDF; requires Tesseract + a vision model (Ollama/Qwen3-VL) configured to extract fields."
        ground_truth[filename] = entry


# --------------------------------------------------------------------------
# PDF invoice generation (digital, template A)
# --------------------------------------------------------------------------

def _po_amount(po_number: str) -> float:
    for num, _vendor, _date, items in PURCHASE_ORDERS:
        if num == po_number:
            return round(sum(q * p for _n, _d, q, p in items), 2)
    raise KeyError(po_number)


def render_template_a(
    path: Path,
    *,
    vendor_name: str,
    invoice_number: str,
    invoice_date: str | None,
    po_number: str | None,
    line_items: list[tuple[str, str, float, float, float]],
    subtotal: float,
    tax: float,
    total: float,
) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 50

    def line(text: str, size: int = 11, dy: int = 18) -> None:
        nonlocal y
        c.setFont("Helvetica", size)
        c.drawString(50, y, text)
        y -= dy

    line(vendor_name, size=14, dy=24)
    line("INVOICE", size=13, dy=24)
    line(f"Invoice Number: {invoice_number}")
    if invoice_date:
        line(f"Invoice Date: {invoice_date}")
    if po_number:
        line(f"PO Number: {po_number}")
    y -= 10
    line("Bill To: Client Corporation Pvt Ltd")
    y -= 10

    line("Item | Description | Qty | Unit Price | Amount", size=10, dy=16)
    for item_name, description, qty, unit_price, amount in line_items:
        line(f"{item_name} | {description} | {qty} | {unit_price:.2f} | {amount:.2f}", size=10, dy=16)

    y -= 10
    line(f"Subtotal: {subtotal:.2f}")
    line(f"Tax: {tax:.2f}")
    line(f"Total: {total:.2f}")

    c.save()


def render_template_b(
    path: Path,
    *,
    vendor_name: str,
    invoice_number: str,
    invoice_date: str | None,
    po_number: str | None,
    line_items: list[tuple[str, str, float, float, float]],
    subtotal: float,
    tax: float,
    total: float,
) -> None:
    """An alternate layout: different labels, different ordering, different
    tax/total wording -- exercises the parser's label-alternative handling.
    """
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 50

    def line(text: str, size: int = 11, dy: int = 18) -> None:
        nonlocal y
        c.setFont("Helvetica", size)
        c.drawString(50, y, text)
        y -= dy

    line("TAX INVOICE", size=13, dy=24)
    line(f"From: {vendor_name}", size=12, dy=20)
    line(f"Inv No: {invoice_number}")
    if invoice_date:
        line(f"Date: {invoice_date}")
    if po_number:
        line(f"Purchase Order Number: {po_number}")
    y -= 10
    line("Ship To: Client Corporation Pvt Ltd, Bengaluru")
    y -= 10

    line("Item | Description | Qty | Unit Price | Amount", size=10, dy=16)
    for item_name, description, qty, unit_price, amount in line_items:
        line(f"{item_name} | {description} | {qty} | {unit_price:.2f} | {amount:.2f}", size=10, dy=16)

    y -= 10
    line(f"Sub Total: {subtotal:.2f}")
    line(f"GST: {tax:.2f}")
    line(f"Grand Total: {total:.2f}")

    c.save()


def render_scanned_invoice(path: Path, *, vendor_name: str, invoice_number: str, po_number: str) -> None:
    """Builds a PDF whose ONLY content is a rendered image (no text layer),
    simulating a scanned paper invoice. Exercises the OCR + vision fallback
    path (and, when Tesseract/Ollama aren't configured, the graceful
    REVIEW-on-extraction-failure path).
    """
    img = Image.new("RGB", (1240, 1754), color="white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 28)
    except OSError:
        font = ImageFont.load_default()

    lines = [
        vendor_name,
        "INVOICE (SCANNED COPY)",
        f"Invoice Number: {invoice_number}",
        "Invoice Date: 2026-01-20",
        f"PO Number: {po_number}",
        "",
        "Item: A4 Paper Ream   Qty: 100   Unit Price: 250.00   Amount: 25000.00",
        "",
        "Subtotal: 25000.00",
        "Tax: 0.00",
        "Total: 25000.00",
    ]
    y = 80
    for text in lines:
        draw.text((80, y), text, fill="black", font=font)
        y += 50

    img_path = path.with_suffix(".png")
    img.save(img_path)

    doc = fitz.open()
    page = doc.new_page(width=img.width, height=img.height)
    page.insert_image(page.rect, filename=str(img_path))
    doc.save(str(path))
    doc.close()
    img_path.unlink(missing_ok=True)


# --------------------------------------------------------------------------
# Build every invoice in the test set
# --------------------------------------------------------------------------

def build_invoices() -> dict:
    INVOICES_DIR.mkdir(parents=True, exist_ok=True)
    ground_truth: dict[str, dict] = {}

    # 1. exact_match -> APPROVE
    render_template_a(
        INVOICES_DIR / "exact_match.pdf",
        vendor_name="Apex Office Supplies Pvt Ltd",
        invoice_number="INV-2001",
        invoice_date="2026-01-15",
        po_number="PO2001",
        line_items=[("A4 Paper Ream", "A4 Paper Ream 500 sheets", 100, 250.00, 25000.00)],
        subtotal=25000.00, tax=0.00, total=25000.00,
    )
    ground_truth["exact_match.pdf"] = {"expected_decision": "APPROVE", "notes": "Exact PO number, vendor, quantity, price, amount match."}

    # 2. price_mismatch -> REVIEW (mirrors the worked example in the spec: PO2009, 120000 vs 130000)
    render_template_a(
        INVOICES_DIR / "price_mismatch.pdf",
        vendor_name="BrightTech Solutions Pvt Ltd",
        invoice_number="INV-2002",
        invoice_date="2026-01-16",
        po_number="PO2009",
        line_items=[("Server Rack Unit", "42U Server Rack Unit", 3, 43333.33, 130000.00)],
        subtotal=130000.00, tax=0.00, total=130000.00,
    )
    ground_truth["price_mismatch.pdf"] = {"expected_decision": "REVIEW", "notes": "Unit price and total exceed PO2009 amount beyond the 2% vendor tolerance."}

    # 3. quantity_mismatch -> REVIEW (overbilled quantity vs PO2003's 12 units)
    render_template_a(
        INVOICES_DIR / "quantity_mismatch.pdf",
        vendor_name="CloudNine Services Pvt Ltd",
        invoice_number="INV-2003",
        invoice_date="2026-01-17",
        po_number="PO2003",
        line_items=[("Cloud Storage Subscription", "Cloud Storage Subscription - Monthly", 15, 2000.00, 30000.00)],
        subtotal=30000.00, tax=0.00, total=30000.00,
    )
    ground_truth["quantity_mismatch.pdf"] = {"expected_decision": "REVIEW", "notes": "Invoice bills 15 units against a PO2003 quantity of 12 (overbilling)."}

    # 4. missing_po -> REVIEW: level-3 candidate matching unambiguously finds PO2004 (vendor + item + qty +
    # price + amount all match), but a missing PO reference is never auto-approved from candidate scoring
    # alone -- it always needs a human to confirm the inferred PO (mirrors the real dataset's R002 policy).
    render_template_a(
        INVOICES_DIR / "missing_po.pdf",
        vendor_name="Apex Office Supplies Pvt Ltd",
        invoice_number="INV-2004",
        invoice_date="2026-01-18",
        po_number=None,
        line_items=[("Whiteboard Markers Box", "Whiteboard Markers Box", 50, 150.00, 7500.00)],
        subtotal=7500.00, tax=0.00, total=7500.00,
    )
    ground_truth["missing_po.pdf"] = {"expected_decision": "REVIEW", "notes": "No PO number on invoice; unambiguously matched to PO2004 via vendor/item/amount similarity, but still routed to REVIEW for human confirmation since the PO reference itself was missing."}

    # 5. tax_total_mismatch -> REVIEW (arithmetic inconsistency: 48000 + 4800 != 50000)
    render_template_a(
        INVOICES_DIR / "tax_total_mismatch.pdf",
        vendor_name="CloudNine Services Pvt Ltd",
        invoice_number="INV-2005",
        invoice_date="2026-01-19",
        po_number="PO2006",
        line_items=[("Onsite Support Hours", "Onsite Technical Support Hours", 40, 1200.00, 48000.00)],
        subtotal=48000.00, tax=4800.00, total=50000.00,
    )
    ground_truth["tax_total_mismatch.pdf"] = {"expected_decision": "REVIEW", "notes": "Stated total (50000) does not equal subtotal+tax (52800); arithmetic check fails."}

    # 6. missing_date -> REVIEW (required field missing) even though amounts match PO2007 exactly
    render_template_a(
        INVOICES_DIR / "missing_date.pdf",
        vendor_name="Delta Facilities Management",
        invoice_number="INV-2006",
        invoice_date=None,
        po_number="PO2007",
        line_items=[("Facility Maintenance Contract", "Facility Maintenance Contract - Quarterly", 1, 65000.00, 65000.00)],
        subtotal=65000.00, tax=0.00, total=65000.00,
    )
    ground_truth["missing_date.pdf"] = {"expected_decision": "REVIEW", "notes": "Invoice date missing; required-fields check fails even though amounts reconcile."}

    # 7-9. split_invoice_{1,2,3} against PO2010 (100 man-days @ 1000 = 100000) -> APPROVE_PARTIAL x3
    # Must be processed in this numeric order for the ledger's cumulative math to work as documented.
    render_template_a(
        INVOICES_DIR / "split_invoice_1.pdf",
        vendor_name="Delta Facilities Management",
        invoice_number="INV-2007-A",
        invoice_date="2026-02-01",
        po_number="PO2010",
        line_items=[("Housekeeping Service - Man Days", "Housekeeping Service - Man Days", 40, 1000.00, 40000.00)],
        subtotal=40000.00, tax=0.00, total=40000.00,
    )
    ground_truth["split_invoice_1.pdf"] = {"expected_decision": "APPROVE_PARTIAL", "notes": "First of 3 split invoices against PO2010 (40/100 man-days)."}

    render_template_a(
        INVOICES_DIR / "split_invoice_2.pdf",
        vendor_name="Delta Facilities Management",
        invoice_number="INV-2007-B",
        invoice_date="2026-02-08",
        po_number="PO2010",
        line_items=[("Housekeeping Service - Man Days", "Housekeeping Service - Man Days", 30, 1000.00, 30000.00)],
        subtotal=30000.00, tax=0.00, total=30000.00,
    )
    ground_truth["split_invoice_2.pdf"] = {"expected_decision": "APPROVE_PARTIAL", "notes": "Second of 3 split invoices against PO2010 (cumulative 70/100 man-days)."}

    render_template_a(
        INVOICES_DIR / "split_invoice_3.pdf",
        vendor_name="Delta Facilities Management",
        invoice_number="INV-2007-C",
        invoice_date="2026-02-15",
        po_number="PO2010",
        line_items=[("Housekeeping Service - Man Days", "Housekeeping Service - Man Days", 30, 1000.00, 30000.00)],
        subtotal=30000.00, tax=0.00, total=30000.00,
    )
    ground_truth["split_invoice_3.pdf"] = {"expected_decision": "APPROVE_PARTIAL", "notes": "Third of 3 split invoices against PO2010; cumulative reaches 100/100 man-days."}

    # 10. scanned_invoice -> REVIEW when OCR/vision aren't configured locally (environment-dependent)
    render_scanned_invoice(
        INVOICES_DIR / "scanned_invoice.pdf",
        vendor_name="Apex Office Supplies Pvt Ltd",
        invoice_number="INV-2008",
        po_number="PO2001",
    )
    ground_truth["scanned_invoice.pdf"] = {
        "expected_decision": "REVIEW",
        "environment_dependent": True,
        "notes": "Image-only PDF (no text layer). Expects REVIEW ('could not be reliably extracted') when Tesseract/Ollama are not configured; may extract successfully (and APPROVE) once they are.",
    }

    # 11. unapproved_vendor -> REJECT
    render_template_a(
        INVOICES_DIR / "unapproved_vendor.pdf",
        vendor_name="Unknown Vendor",
        invoice_number="INV-2009",
        invoice_date="2026-01-20",
        po_number="PO8888",
        line_items=[("Miscellaneous Goods", "Miscellaneous Goods", 1, 10000.00, 10000.00)],
        subtotal=10000.00, tax=0.00, total=10000.00,
    )
    ground_truth["unapproved_vendor.pdf"] = {"expected_decision": "REJECT", "notes": "Vendor 'Unknown Vendor' (V005) has status NOT_APPROVED in the vendor master."}

    # 12. different_layout -> APPROVE (template B: different labels/order, same PO2008 numbers)
    render_template_b(
        INVOICES_DIR / "different_layout.pdf",
        vendor_name="Apex Office Supplies Pvt Ltd",
        invoice_number="INV-2010",
        invoice_date="15-Jan-2026",
        po_number="PO2008",
        line_items=[("Ergonomic Chair", "Ergonomic Office Chair", 10, 8000.00, 80000.00)],
        subtotal=80000.00, tax=0.00, total=80000.00,
    )
    ground_truth["different_layout.pdf"] = {"expected_decision": "APPROVE", "notes": "Alternate invoice layout/label wording (template B) for the same underlying PO2008 match."}

    # 13. po_number_format_variant -> APPROVE (PO-2005 normalizes to PO2005, exact level-1 match)
    render_template_a(
        INVOICES_DIR / "po_number_format_variant.pdf",
        vendor_name="BrightTech Solutions Pvt Ltd",
        invoice_number="INV-2011",
        invoice_date="2026-01-21",
        po_number="PO-2005",
        line_items=[("Network Switch 24-port", "Managed Network Switch 24-port", 5, 12000.00, 60000.00)],
        subtotal=60000.00, tax=0.00, total=60000.00,
    )
    ground_truth["po_number_format_variant.pdf"] = {"expected_decision": "APPROVE", "notes": "PO written as 'PO-2005'; normalizes to 'PO2005' for an exact level-1 match."}

    # 14. malformed_po -> APPROVE (level-2 digits-only match: "Ref: 2011" -> PO2011)
    render_template_a(
        INVOICES_DIR / "malformed_po.pdf",
        vendor_name="BrightTech Solutions Pvt Ltd",
        invoice_number="INV-2012",
        invoice_date="2026-01-22",
        po_number="Ref: 2011",
        line_items=[("USB-C Hub", "USB-C Hub 7-in-1", 25, 800.00, 20000.00)],
        subtotal=20000.00, tax=0.00, total=20000.00,
    )
    ground_truth["malformed_po.pdf"] = {"expected_decision": "APPROVE", "notes": "PO reference given only as 'Ref: 2011' (malformed); matched to PO2011 via digits-only normalization (level 2)."}

    # 15. bundled_line_items -> REVIEW (invoice bundles PO2012's two line items into one; unit-price/line detail can't be reconciled automatically)
    render_template_a(
        INVOICES_DIR / "bundled_line_items.pdf",
        vendor_name="Apex Office Supplies Pvt Ltd",
        invoice_number="INV-2013",
        invoice_date="2026-01-23",
        po_number="PO2012",
        line_items=[("Stationery Bundle", "Stationery Bundle - Pens Box and Notebooks Pack", 1, 9000.00, 9000.00)],
        subtotal=9000.00, tax=0.00, total=9000.00,
    )
    ground_truth["bundled_line_items.pdf"] = {
        "expected_decision": "REVIEW",
        "notes": "Invoice bundles PO2012's two separate line items (Pens Box + Notebooks Pack) into a single line; unit-price reconciliation fails at line-item level and is correctly routed to REVIEW rather than silently approved.",
    }

    for entry in ground_truth.values():
        entry.setdefault("source", "synthetic (this script)")

    return ground_truth


def main() -> None:
    build_vendor_xlsx()
    build_po_xlsx()
    ground_truth = build_invoices()
    copy_real_invoices_and_merge_ground_truth(ground_truth)
    with open(INVOICES_DIR / "ground_truth.json", "w", encoding="utf-8") as f:
        json.dump(ground_truth, f, indent=2)
    print(f"Wrote {DATA_DIR / 'vendors.xlsx'}")
    print(f"Wrote {DATA_DIR / 'purchase_orders.xlsx'}")
    print(f"Wrote {len(ground_truth)} invoice PDFs to {INVOICES_DIR} (real + synthetic)")
    print(f"Wrote {INVOICES_DIR / 'ground_truth.json'}")


if __name__ == "__main__":
    main()
